# -*- coding: utf-8 -*-
"""
51Talk Big Team 01 - Hourly Performance Monitor (Real-Time Live Cloud Engine)
Dynamically links to the active daily tasks sheet on DingTalk, downloads live updates,
calculates exact KPIs and team breakdowns, and sends an executive report to Saber Hussien.
"""
import os
import sys
import json
import subprocess
from datetime import datetime
from pathlib import Path
import openpyxl

# Enforce UTF-8 stdout
sys.stdout.reconfigure(encoding='utf-8')

# Target Recipient: Saber Hussien (Personal 1-on-1 Chat)
RECEIVER_OPEN_DINGTALK_ID = "DMmZ0gsytpviSWm51piSq3GstTd6V9Voffs"

# Node & CLI Candidates
NODE_BIN_CANDIDATES = [
    r"C:\Program Files\nodejs\node.exe",
    r"C:\Program Files (x86)\nodejs\node.exe",
    "node"
]
node_bin = next((p for p in NODE_BIN_CANDIDATES if os.path.exists(p)), "node")
dws_js = r"C:\Users\husse\AppData\Roaming\npm\node_modules\dingtalk-workspace-cli\bin\dws.js"
runner_cmd = [node_bin, dws_js] if os.path.exists(node_bin) and os.path.exists(dws_js) else ["dws"]

TEMP_LIVE_PATH = r"C:\Users\husse\.gemini\antigravity-ide\scratch\October_Daily_Tasks_live_hourly.xlsx"
URL_FILE_PATH = Path(r"D:\Lens\Dashboard\latest_online_sheet_url.txt")
ACTIVE_NODE_JSON = Path(r"D:\Lens\Dashboard\active_daily_task_node.json")

# 21 Representatives Roster across 5 Teams
REP_ROSTER = {
    'ashraqatal': 'Team 01', 'juliamonir01': 'Team 01', 'mahmoud04': 'Team 01', 'negma': 'Team 01', 'nohayoussry': 'Team 01',
    'abdelrahmannasef': 'Team 05', 'ehabzaky01': 'Team 05', 'ibrahimismaiel': 'Team 05', 'khaledgonam': 'Team 05', 'omarmoneb': 'Team 05', 'samira01': 'Team 05',
    'abdelrhmanshehata': 'Team 10', 'ahmedshoukry': 'Team 10', 'mahmoudkhamis': 'Team 10',
    'amrsafwat': 'Team 13', 'hayamhassan': 'Team 13', 'marwaahmed': 'Team 13', 'mohamedha': 'Team 13',
    'adhmgadallah': 'Team 30', 'alihesham01': 'Team 30', 'titooooo': 'Team 30'
}


def auto_detect_live_node():
    """Detects active online sheet node for today, with dynamic cloud fallback."""
    # 1. Try reading stored configuration
    if URL_FILE_PATH.exists():
        try:
            u = URL_FILE_PATH.read_text(encoding="utf-8-sig").strip()
            if "nodes/" in u:
                nid = u.split("nodes/")[1].split("?")[0].strip()
                return nid, u
        except Exception:
            pass

    # 2. Query DingTalk Drive for newest October Daily Tasks file
    try:
        cmd = runner_cmd + ["drive", "search", "--query", "October Daily Tasks", "-y"]
        res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=15)
        if res.returncode == 0:
            data = json.loads(res.stdout)
            docs = data.get("doc_results", {}).get("documents", [])
            if docs:
                latest_doc = docs[0]
                nid = latest_doc.get("nodeId")
                doc_url = latest_doc.get("docUrl", f"https://alidocs.dingtalk.com/i/nodes/{nid}")
                if "alidocs.dingtalk.com" not in doc_url:
                    doc_url = f"https://alidocs.dingtalk.com/i/nodes/{nid}"
                # Cache for next runs
                try:
                    URL_FILE_PATH.write_text(doc_url, encoding="utf-8")
                except Exception:
                    pass
                return nid, doc_url
    except Exception as e:
        print(f"[WARN] Drive search fallback: {e}")

    # 3. Default fallback node for today
    fallback_nid = "dxXB52LJqnajQAlEH2krM9LL8qjMp697"
    fallback_url = f"https://alidocs.dingtalk.com/i/nodes/{fallback_nid}"
    return fallback_nid, fallback_url


def download_live_workbook(node_id):
    """Downloads real-time sheet from DingTalk Drive node."""
    print(f"[SYNC] Downloading live rep updates from DingTalk Drive (Node: {node_id})...", flush=True)
    cmd = runner_cmd + [
        "drive", "download",
        "--node", node_id,
        "--output", TEMP_LIVE_PATH,
        "--overwrite", "-y"
    ]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30)
        if res.returncode == 0 and os.path.exists(TEMP_LIVE_PATH) and os.path.getsize(TEMP_LIVE_PATH) > 10000:
            print(f"[SYNC] Successfully downloaded latest live cloud sheet ({os.path.getsize(TEMP_LIVE_PATH)} bytes).", flush=True)
            return TEMP_LIVE_PATH
    except Exception as e:
        print(f"[WARN] Live cloud download error: {e}. Checking local copies...", flush=True)

    # Local fallbacks
    fallbacks = [
        r"C:\Users\husse\Desktop\October Daily Tasks.xlsx",
        r"D:\Lens\Dashboard\October Daily Tasks.xlsx",
        r"D:\Lens\Dashboard\Dashboard_Input_Files\October Daily Tracker\October Daily Tasks.xlsx"
    ]
    for fb in fallbacks:
        if os.path.exists(fb) and os.path.getsize(fb) > 10000:
            print(f"[SYNC] Using local fallback workbook: {fb}", flush=True)
            return fb

    raise FileNotFoundError("Could not find or download October Daily Tasks workbook.")


def calculate_live_progress(wb_path):
    """Parses all 21 rep sheets in the workbook and calculates exact live metrics."""
    wb = openpyxl.load_workbook(wb_path, data_only=False)
    teams = ["Team 01", "Team 05", "Team 10", "Team 13", "Team 30"]
    team_stats = {t: {"total": 0, "completed": 0, "pending": 0} for t in teams}
    rep_details = []
    tier_counts = {"P1": 0, "P2": 0, "P3": 0, "P4": 0, "P5": 0}

    total_assigned = 0
    total_completed = 0

    for rep, team in REP_ROSTER.items():
        sheet_name = f"{team} - {rep}"
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rep_assigned = 0
        rep_done = 0

        for r in range(6, ws.max_row + 1):
            sid = ws.cell(r, 4).value  # Col D: Student ID
            if not sid:
                continue
            rep_assigned += 1

            tier_val = str(ws.cell(r, 1).value or '')  # Col A: Priority Tier
            if "P1" in tier_val or "Round 1" in tier_val: tier_counts["P1"] += 1
            elif "P2" in tier_val or "Due Today" in tier_val: tier_counts["P2"] += 1
            elif "P3" in tier_val or "Unfixed" in tier_val: tier_counts["P3"] += 1
            elif "P4" in tier_val or "0-Class" in tier_val: tier_counts["P4"] += 1
            elif "P5" in tier_val or "Outside" in tier_val: tier_counts["P5"] += 1

            status_val = str(ws.cell(r, 10).value or '').strip()  # Col J: Status
            feedback_val = str(ws.cell(r, 11).value or '').strip()  # Col K: Feedback

            is_done = False
            if status_val.lower() in ['true', '1', 'done', 'done ✓', 'checked', 'yes'] or 'done' in status_val.lower():
                is_done = True
            elif len(feedback_val) > 1 and status_val.lower() not in ['false', '0']:
                is_done = True

            if is_done:
                rep_done += 1

        team_stats[team]["total"] += rep_assigned
        team_stats[team]["completed"] += rep_done
        total_assigned += rep_assigned
        total_completed += rep_done

        rate = round(rep_done / rep_assigned * 100, 1) if rep_assigned > 0 else 0.0
        rep_details.append({
            "rep": rep,
            "team": team,
            "assigned": rep_assigned,
            "completed": rep_done,
            "pending": rep_assigned - rep_done,
            "rate": rate
        })

    for tm in teams:
        team_stats[tm]["pending"] = team_stats[tm]["total"] - team_stats[tm]["completed"]

    total_pending = total_assigned - total_completed
    completion_rate = round(total_completed / total_assigned * 100, 1) if total_assigned > 0 else 0.0

    # Top performers (reps with at least 1 completed task, sorted descending)
    active_reps = [r for r in rep_details if r["completed"] > 0]
    active_reps.sort(key=lambda x: (x["completed"], x["rate"]), reverse=True)

    return {
        "total_assigned": total_assigned,
        "total_completed": total_completed,
        "total_pending": total_pending,
        "completion_rate": completion_rate,
        "team_stats": team_stats,
        "tier_counts": tier_counts,
        "active_reps": active_reps
    }


def main():
    print("=" * 68, flush=True)
    print(" 51Talk Big Team 01 - Hourly Performance Monitor (Live Cloud Engine)", flush=True)
    print(" Target Recipient: Saber Hussien (EGSS-Hussienmo)", flush=True)
    print("=" * 68, flush=True)

    node_id, online_url = auto_detect_live_node()
    wb_path = download_live_workbook(node_id)
    stats = calculate_live_progress(wb_path)

    current_time_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    teams = ["Team 01", "Team 05", "Team 10", "Team 13", "Team 30"]

    team_lines = []
    for tm in teams:
        st = stats["team_stats"][tm]
        pct = round((st["completed"] / st["total"]) * 100, 1) if st["total"] > 0 else 0.0
        team_lines.append(f"• {tm}: الإجمالي {st['total']} | المنجز {st['completed']} ({pct}%) | المتبقي {st['pending']}")

    team_text = "\n".join(team_lines)

    # Top performers block
    top_lines = []
    for r in stats["active_reps"][:6]:
        medal = " 🏆" if r["rate"] >= 100.0 else (" 🚀" if r["rate"] >= 50.0 else "")
        top_lines.append(f"• {r['team']} - {r['rep']}: أنجز {r['completed']} من {r['assigned']} ({r['rate']}%){medal}")

    top_text = "\n".join(top_lines) if top_lines else "• جاري بدء تحديثات المتابعة من الممثلين..."

    report_text = f"""📊 51Talk Big Team 01 - التقرير اللحظي الدوري (Hourly Progress)
───────────────────────────────────
👤 المستلم: Saber Hussien (تقرير تنفيذي خاص)
⏰ التوقيت: {current_time_str} | النطاق: Big Team 01 (21 SS Agents)

📈 ملخص الإنجاز اللحظي (Executive KPIs):
• 🎯 إجمالي مهام اليوم المستهدفة: {stats['total_assigned']} عميل
• ✅ المهام المنجزة: {stats['total_completed']} عميل ({stats['completion_rate']}%)
• ⏳ المهام الجاري متابعتها: {stats['total_pending']} عميل

🏢 موقف إنجاز الفرق (Team Completion Breakdown):
{team_text}

🌟 مميزو الإنجاز اللحظي (Top Active Performers):
{top_text}

📌 تفاصيل الأولويات التشغيلية:
• 🔴 P1 (SOP Round 1): {stats['tier_counts']['P1']} عميل
• ⏰ P2 (Due Today Tasks): {stats['tier_counts']['P2']} عميل
• 🎓 P3 (Unfixed Teachers - M0, M1, M2): {stats['tier_counts']['P3']} عميل
• 📉 P4 (0-Class Consumption - M0, M1, M2): {stats['tier_counts']['P4']} عميل
• 🔄 P5 (Outside Pool & 0 Pts): {stats['tier_counts']['P5']} عميل

🌐 رابط متابعة الشيت أونلاين على دينج توك:
👉 {online_url}
───────────────────────────────────
(يتم تحديث التقرير الدوري آلياً كل ساعة لحساب Saber Hussien حصرياً)"""

    # Sanitize any hidden zero-width unicode characters
    clean_report = report_text.replace('\u200d', '').replace('\u200c', '').replace('\u200b', '').replace('\ufeff', '')

    print(f"\n[METRICS] Total: {stats['total_assigned']} | Completed: {stats['total_completed']} ({stats['completion_rate']}%) | Pending: {stats['total_pending']}", flush=True)
    for tm in teams:
        st = stats["team_stats"][tm]
        r = round((st['completed']/st['total'])*100, 1) if st['total'] > 0 else 0
        print(f"   {tm}: Total {st['total']} | Done {st['completed']} | Pend {st['pending']} ({r}%)", flush=True)

    print("\nBroadcasting Full Expanded Report to Saber Hussien...", flush=True)

    try:
        cmd = runner_cmd + [
            "chat", "+messages-send",
            "--as", "user",
            "--open-dingtalk-id", RECEIVER_OPEN_DINGTALK_ID,
            "--text", clean_report,
            "--ai-tag=false",
            "-y"
        ]
        res = subprocess.run(cmd, input="y\n", capture_output=True, text=True, encoding='utf-8', errors='replace')
        if res.returncode == 0:
            print("[SUCCESS] Full text hourly progress report delivered to Saber Hussien successfully!", flush=True)
        else:
            print(f"[STATUS] ReturnCode: {res.returncode}", flush=True)
            if res.stdout: print(f"[STDOUT] {res.stdout.strip()}", flush=True)
            if res.stderr: print(f"[STDERR] {res.stderr.strip()}", flush=True)
    except Exception as e:
        print(f"[ERROR] Exception during send: {e}", flush=True)


if __name__ == "__main__":
    main()
